"""Genera cinco instrumentos institucionales a partir de las plantillas DOCX locales."""

import html
import os
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).parent
VERTICAL = ROOT / "Plantilla_Vertical.docx"
MANUAL = ROOT / "Plantilla_Manual.docx"
NAVY = RGBColor(9, 34, 64)
BLUE = RGBColor(52, 118, 177)
RED = RGBColor(218, 24, 30)
GRAY = RGBColor(87, 96, 106)


def style_document(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = NAVY
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12
    for key, size, before, after in [
        ("Title", 20, 2, 9),
        ("Heading 1", 12.5, 13, 6),
        ("Heading 2", 10.5, 9, 4),
    ]:
        s = doc.styles[key]
        s.font.name = "Arial"
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = NAVY if key == "Title" else BLUE
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True
    if "InstrumentoOpcion" not in [s.name for s in doc.styles]:
        opt = doc.styles.add_style("InstrumentoOpcion", WD_STYLE_TYPE.PARAGRAPH)
        opt.base_style = normal
    opt = doc.styles["InstrumentoOpcion"]
    opt.font.name = "Arial"
    opt.font.size = Pt(9.2)
    opt.paragraph_format.left_indent = Cm(0.55)
    opt.paragraph_format.space_after = Pt(1)


def base(template):
    doc = Document(template)
    # Manual: conservar el primer párrafo con la imagen y los cuadros de portada.
    # Vertical: conservar los elementos de marca del encabezado y pie, no sus
    # párrafos de cuerpo vacíos.
    keep_cover = template == MANUAL
    for p in list(doc.paragraphs)[1 if keep_cover else 0:]:
        p._element.getparent().remove(p._element)
    style_document(doc)
    return doc


def paragraph(doc, text="", bold_prefix=None, style=None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def bullet(doc, text, number):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.35)
    run = p.add_run(f"Paso {number}. ")
    run.bold = True
    run.font.color.rgb = BLUE
    p.add_run(text)
    return p


def heading(doc, text, level=1):
    return doc.add_heading(text, level)


def shade(cell, fill="EAF1F8"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tcPr.insert_element_before(shd, "w:noWrap", "w:tcMar", "w:textDirection",
                               "w:tcFitText", "w:vAlign", "w:hideMark")


def table_borders(table):
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), "D6DEE8")
        borders.append(edge)
    table._tbl.tblPr.insert_element_before(
        borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
        "w:tblCaption", "w:tblDescription")


def table_widths(table, widths):
    """Fija también tblGrid/tblW: tcW solo permite que Word ensanche la tabla."""
    table.autofit = False
    for col, width in zip(table.columns, widths):
        col.width = width
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width
    tblW = table._tbl.tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        table._tbl.tblPr.insert(0, tblW)
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(sum(w.twips for w in widths)))


def prevent_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def set_cell(cell, text, *, bold=False):
    cell.text = text
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(2)
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(8.5)
            run.bold = bold


def title_and_meta(doc, title, kind, day, points, duration, diagnostic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("CEC EPN  /  DIPLOMADO EN BIG DATA E INTELIGENCIA ARTIFICIAL  /  MÓDULO 3")
    r.bold = True
    r.font.size = Pt(8)
    r.font.color.rgb = BLUE
    doc.add_paragraph(title, style="Title")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(10)
    run = subtitle.add_run(f"{kind} · Día {day} · {points} · {duration}")
    run.bold = True
    run.font.color.rgb = RED
    run.font.size = Pt(10)
    t = doc.add_table(rows=2, cols=2)
    t.autofit = False
    for row in t.rows:
        prevent_split(row)
        for c in row.cells:
            c.width = Cm(7.5)
    table_widths(t, [Cm(7.5), Cm(7.5)])
    for cell, text in zip([t.cell(0,0),t.cell(0,1),t.cell(1,0),t.cell(1,1)],
                          ["Participante/equipo: ___________________", "Fecha: ____ / ____ / 2026",
                           ("Grupo/horario: _______________________" if diagnostic else "Dataset elegido: _______________________"),
                           "Instructor: ______________________"]):
        set_cell(cell,text)
        shade(cell,"F1F5F9")
    doc.add_paragraph()


def sections(doc, description, objective, expected, work):
    heading(doc, "Descripción de la actividad")
    paragraph(doc, description)
    heading(doc, "Objetivo de aprendizaje")
    paragraph(doc, objective)
    heading(doc, "Resultado esperado")
    paragraph(doc, expected)
    heading(doc, "Cómo se trabajará")
    for number, item in enumerate(work, 1):
        bullet(doc, item, number)


def rubric(doc, rows, total, note):
    heading(doc, f"Rúbrica de evaluación — {total} puntos")
    paragraph(doc, "Por cada criterio marque Logrado (puntaje completo), En desarrollo (mitad) o No logrado (0). Se admiten medios puntos. Registre el puntaje real y una observación específica cuando falte evidencia.")
    t = doc.add_table(rows=1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(t)
    widths = [Cm(3.0), Cm(10.3), Cm(1.7)]
    table_widths(t, widths)
    for i,(cell,text) in enumerate(zip(t.rows[0].cells,["Criterio","Niveles de desempeño y evidencia observable","Puntos"])):
        cell.width=widths[i];set_cell(cell,text,bold=True);shade(cell,"DDEAF5")
        cell.paragraphs[0].paragraph_format.keep_with_next = True
    for criterion,maxpt,full,partial,missing in rows:
        cells=t.add_row().cells
        prevent_split(t.rows[-1])
        for cell,width in zip(cells,widths):cell.width=width
        set_cell(cells[0],criterion,bold=True)
        desc=cells[1];desc.text=""
        for label,score,detail in [("Logrado",str(maxpt),full),
                                   ("En desarrollo",str(maxpt/2).rstrip('0').rstrip('.') if isinstance(maxpt/2,float) else str(maxpt/2),partial),
                                   ("No logrado","0",missing)]:
            p=desc.paragraphs[0] if not desc.paragraphs[0].text else desc.add_paragraph()
            p.paragraph_format.space_after=Pt(2)
            run=p.add_run(f"{label} ({score}): ");run.bold=True;run.font.size=Pt(8.3)
            run=p.add_run(detail);run.font.size=Pt(8.3)
        set_cell(cells[2],f"___ / {maxpt}")
    total_cells=t.add_row().cells
    prevent_split(t.rows[-1])
    set_cell(total_cells[0],"TOTAL",bold=True)
    set_cell(total_cells[1],note)
    set_cell(total_cells[2],f"___ / {total}",bold=True)
    for c in total_cells:shade(c,"EAF1F8")
    table_widths(t, widths)
    paragraph(doc,"Observaciones y retroalimentación: __________________________________________________________")
    paragraph(doc,"____________________________________________________________________________________")


QUESTIONS = [
    ("Python", "Si x = 3 y luego se ejecuta x = x + 2, ¿cuál es el valor final de x?", ["1", "3", "5", "Se produce un error"], "C", "La segunda asignación suma 2 al valor anterior: 3 + 2 = 5."),
    ("Python", "En Python, ¿qué devuelve [10, 20, 30][1]?", ["10", "20", "30", "1"], "B", "Los índices empiezan en 0; el índice 1 corresponde a 20."),
    ("Python", "¿Qué devuelve len([2, 4, 6])?", ["2", "3", "6", "Un error"], "B", "La lista tiene tres elementos."),
    ("Python", "¿Qué valores recorre i en for i in range(3)?", ["1, 2, 3", "0, 1, 2", "0, 1, 2, 3", "Solo 3"], "B", "range(3) incluye 0 y se detiene antes de 3."),
    ("NumPy", "Con import numpy as np, ¿qué resulta de np.array([1, 2, 3]) * 2?", ["[1, 2, 3, 1, 2, 3]", "[2, 4, 6]", "[3, 4, 5]", "Un error"], "B", "La multiplicación se aplica a cada elemento del arreglo."),
    ("NumPy", "¿Cuál es la forma (shape) de np.array([[1, 2], [3, 4]])?", ["(4,)", "(2, 2)", "(1, 4)", "(2,)"], "B", "El arreglo tiene dos filas y dos columnas."),
    ("NumPy", "¿Cuánto vale np.mean(np.array([2, 4, 6]))?", ["2", "3", "4", "6"], "C", "La media es (2 + 4 + 6) / 3 = 4."),
    ("NumPy", "Si a = np.array([1, 4, 7]), ¿qué devuelve a[a > 3]?", ["[1]", "[1, 4, 7]", "[4, 7]", "[3, 4, 7]"], "C", "La máscara booleana conserva solo los valores mayores que 3."),
    ("Limpieza de datos", "¿Para qué sirve df.isna().sum() en una tabla de pandas?", ["Contar valores ausentes por columna", "Eliminar todas las columnas", "Calcular promedios", "Convertir texto en números"], "A", "isna identifica ausencias y sum cuenta las marcas True por columna."),
    ("Limpieza de datos", "¿Qué informa df.duplicated().sum() por defecto?", ["El número de columnas numéricas", "El número de filas marcadas como duplicadas de una fila previa", "El número de valores vacíos", "La media de las filas repetidas"], "B", "duplicated marca filas repetidas y sum cuenta esas marcas."),
    ("Limpieza de datos", "Si se rellenarán edades faltantes con una mediana, ¿cuál es el procedimiento correcto?", ["Calcular la mediana usando Train, Validation y Test juntos", "Poner edad cero sin revisar", "Calcular la mediana solo en Train y aplicarla a Validation/Test", "Eliminar siempre todo el dataset"], "C", "La transformación se aprende en Train para no filtrar información de las filas reservadas."),
    ("Limpieza de datos", "Queremos decidir a quién llamar antes de la llamada. ¿Por qué no usar su duración como predictor?", ["Porque siempre es texto", "Porque no existe la clase yes", "Porque todas las llamadas duran igual", "Porque solo se conoce después y causaría fuga de información"], "D", "duration aún no existe en el momento de la decisión."),
    ("Probabilidad", "En un dado justo de seis caras, ¿cuál es la probabilidad de obtener un número par?", ["1/6", "3/6", "4/6", "6/6"], "B", "Hay tres resultados pares entre seis equiprobables."),
    ("Probabilidad", "Si P(A) = 0,30, ¿cuánto vale P(no A)?", ["0,30", "0,50", "0,70", "1,30"], "C", "El complemento es 1 - 0,30 = 0,70."),
    ("Probabilidad", "Si A y B son independientes, P(A)=0,20 y P(B)=0,50, ¿cuánto vale P(A y B)?", ["0,10", "0,20", "0,50", "0,70"], "A", "Para eventos independientes se multiplican 0,20 × 0,50 = 0,10."),
    ("Probabilidad", "De 20 personas con la característica B, 8 también tienen A. ¿Cuál es P(A dado B)?", ["8/100", "20/8", "12/20", "8/20"], "D", "El grupo de referencia son las 20 personas con B; 8/20 = 0,40."),
    ("Estadística", "¿Cuál es la media de 2, 4 y 6?", ["2", "3", "4", "6"], "C", "La media es (2 + 4 + 6) / 3 = 4."),
    ("Estadística", "¿Cuál es la mediana de 2, 4 y 100?", ["2", "4", "35,3", "100"], "B", "El dato central del conjunto ordenado es 4."),
    ("Estadística", "Si un conjunto tiene mayor desviación estándar que otro en la misma escala, ¿qué sugiere?", ["Sus valores están más dispersos alrededor de su media", "Su media es siempre mayor", "Todos sus valores son iguales", "Existe causalidad"], "A", "Una desviación estándar mayor indica más dispersión, sin informar por sí sola la media."),
    ("Estadística", "Una correlación de 0,8 entre dos variables indica, en general:", ["Que una causa necesariamente a la otra", "Que ambas son idénticas", "Que no existe asociación", "Una asociación lineal positiva fuerte, sin demostrar causalidad"], "D", "La correlación describe asociación lineal; no prueba causa."),
]
assert len(QUESTIONS) == 20
assert all(len(q[2]) == 4 and q[3] in 'ABCD' for q in QUESTIONS)
assert {cat:sum(q[0] == cat for q in QUESTIONS) for cat in {q[0] for q in QUESTIONS}} == {
    "Python":4,"NumPy":4,"Limpieza de datos":4,"Probabilidad":4,"Estadística":4}


def diagnostic():
    doc = base(VERTICAL)
    title_and_meta(doc,"Evaluación diagnóstica · Fundamentos para ML","Diagnóstica",1,"20 × 0,5 = 10 puntos diagnósticos","20–25 minutos",diagnostic=True)
    sections(doc,
        "Cuestionario de entrada para identificar conocimientos previos de Python, NumPy, limpieza de datos, probabilidad y estadística. No es un examen aprobatorio ni modifica la nota del módulo.",
        "Reconocer las habilidades que necesitan refuerzo antes de modelar: operaciones básicas, arreglos, datos faltantes y conceptos probabilísticos/estadísticos.",
        "Una hoja con 20 respuestas de opción múltiple. Cada respuesta correcta suma 0,5: máximo 10 puntos diagnósticos, equivalentes a 0 puntos en la calificación de 100 del módulo. El docente identifica áreas de apoyo por bloque de cuatro preguntas.",
        ["Trabaja individualmente y sin ejecutar código; lee cada enunciado y marca una sola alternativa A, B, C o D.",
         "No se descuentan puntos por error. Respuesta en blanco o más de una opción marcada: 0 en esa pregunta.",
         "El docente revisará los cinco bloques y explicará los temas con más respuestas incorrectas antes de avanzar."])
    heading(doc,"Preguntas · cuatro opciones, una correcta")
    current = None
    for i,(cat,q,opts,answer,why) in enumerate(QUESTIONS,1):
        if cat != current:
            heading(doc,f"Bloque {cat} · preguntas {i}–{i+3}",2)
            current=cat
        p=doc.add_paragraph()
        p.paragraph_format.space_before=Pt(6)
        p.paragraph_format.space_after=Pt(2)
        p.paragraph_format.keep_with_next=True
        run=p.add_run(f"{i:02d}. {q}")
        run.bold=True;run.font.color.rgb=NAVY
        for j,opt_text in enumerate(opts):
            option=doc.add_paragraph(style="InstrumentoOpcion")
            option.add_run(f"{chr(65+j)}) {opt_text}")
            option.paragraph_format.keep_with_next = j < 3
        if i in (4,8,12,16):
            # Separar dominios visualmente sin imponer saltos de página.
            doc.add_paragraph().paragraph_format.space_after=Pt(2)
    heading(doc,"Registro diagnóstico (completa el instructor después de revisar)",2)
    paragraph(doc,"Python ___/2 · NumPy ___/2 · Limpieza ___/2 · Probabilidad ___/2 · Estadística ___/2 · TOTAL BRUTO ___/10. Este resultado no se suma a la nota del módulo.")
    return doc


DATASET_OPTIONS = [
    ("A. Regresión", "actividad_economica_reto.csv", "Estimar crecimiento de actividad por fecha, región y sector."),
    ("B. Clasificación", "bank_marketing_uci.csv", "Predecir yes/no antes de llamar; excluir duration."),
    ("C. Clasificación", "telco_churn_fallback.csv", "Identificar Churn=Yes/No en clientes sintéticos."),
    ("D. Clustering", "world_bank_latam_ml.csv", "Agrupar país-años por indicadores, sin etiqueta de grupo."),
    ("E. Regresión temporal", "macro_mensual_demo.csv", "Pronosticar índice de actividad del mes siguiente con rezagos."),
]


def option_table(doc):
    heading(doc,"Una elección, cinco datasets disponibles")
    t=doc.add_table(rows=1,cols=3);table_borders(t);t.autofit=False
    widths=[Cm(2.7),Cm(4.5),Cm(7.8)]
    for i,(c,txt) in enumerate(zip(t.rows[0].cells,["Tipo orientativo","Archivo en datos/","Pregunta inicial"])):
        c.width=widths[i];set_cell(c,txt,bold=True);shade(c)
    for kind,filename,problem in DATASET_OPTIONS:
        row=t.add_row();prevent_split(row)
        for c,width,txt in zip(row.cells,widths,[kind,filename,problem]):c.width=width;set_cell(c,txt)
    table_widths(t,widths)
    paragraph(doc,"La elección se justifica según la respuesta buscada: cantidad numérica → regresión; etiqueta conocida → clasificación; sin grupo correcto conocido → clustering. No convertir identificadores como customerID o pais a números arbitrarios.")


R1 = [
    ("Pregunta y tipo",4,"Pregunta y decisión claras; explica el tipo según y numérica, categórica o ausencia de grupo conocido.","Identifica el tipo, pero la pregunta o el motivo es impreciso.","No define problema o confunde tipo y respuesta."),
    ("Fuente y fila",4,"Archivo local, fuente pública/sintética, periodo, población, unidad de fila y target si aplica.","Identifica archivo y fila, pero omite contexto importante.","No identifica origen ni unidad de observación."),
    ("Auditoría inicial",5,"Tamaño, columnas, tipos, nulos, duplicados, categorías especiales y extremos explicados.","Revisa solo parte de los datos o no interpreta hallazgos.","No revisa datos o imprime filas sin explicación."),
    ("Primer gráfico",3,"Figura legible con interpretación y dos gráficos futuros pertinentes propuestos.","Figura sin lectura o plan de visualización incompleto.","No hay figura útil."),
    ("Plan de evaluación",4,"Partición apropiada, predictor disponible al decidir y baseline o criterio de agrupación.","Hay split propuesto sin justificación o sin referencia.","No hay protocolo o ya usa Test para decidir."),
]
R2 = [
    ("Limpieza",5,"Trata nulos/duplicados/extremos, excluye IDs y variables futuras y cuantifica pérdida de filas.","Limpia, pero sin justificar una decisión importante.","No limpia o introduce target/futuro como predictor."),
    ("Tres gráficos",5,"Tres figuras distintas con título, ejes y conclusión pertinente por figura.","Una figura falta o las interpretaciones son parciales.","No hay exploración útil."),
    ("Particiones",5,"Train/Validation/Test reproducibles según clases, tiempo, país o estabilidad; comprueba tamaños.","Divide pero no justifica ni revisa composición.","No divide o mezcla Test con Train."),
    ("Sin fuga",5,"Categorías relevantes con One-Hot y transformaciones aprendidas solo en Train; aplica al resto sin reajuste.","Algunas transformaciones son correctas pero falta justificar o aislar un paso.","Codifica usando toda la muestra antes del split o trata IDs como magnitudes."),
]
R3 = [
    ("Referencia",4,"Baseline supervisado o referencia de perfiles/particiones en clustering, con propósito.","Muestra referencia sin explicar qué aporta.","No hay comparación pertinente."),
    ("Alternativas",6,"Modelos o particiones adecuados comparados sobre las mismas filas de Validation sin fuga.","Candidatos presentes pero comparación parcial.","Compara conjuntos distintos o solo Train."),
    ("Métricas",4,"Usa MAE/RMSE, matriz/Recall/Precision/F1 o Silhouette/tamaños/perfiles según tipo, con lectura.","Métrica adecuada sin unidad o interpretación.","Métrica incompatible: p. ej., Accuracy de clustering sin etiquetas."),
    ("Elección",4,"Relaciona selección con objetivo, costos/estabilidad y complejidad o variabilidad.","Escoge un número mayor/menor sin explicar su significado.","Elige arbitrariamente o mirando Test."),
    ("Decisión fijada",2,"Registra regla, modelo/k, parámetros, variables y dos límites antes de Test.","Decisión anotada pero incompleta.","No fija decisión o ajusta después de Test."),
]
RF = [
    ("Comprobación reservada",12,"Test se usa una vez tras fijar regla y métricas; clustering revisa perfiles/estabilidad retenidos sin y ni Accuracy.","Hay comprobación, pero no queda claro que la decisión y las transformaciones estuvieran fijadas.","Ajusta con Test, usa métrica incompatible o falta comprobación."),
    ("Interpretación",10,"Compara con baseline/Validation; explica FP/FN o tamaños/perfiles y reconoce ausencia de mejora.","Presenta cifras sin lectura completa de errores o referencia.","Afirma éxito sin evidencia o interpreta mal resultados."),
    ("Reproducibilidad",8,"Notebook único de Tareas 1–3, CSV local, ejecución ordenada y Test sin fuga.","Integra avances pero requiere pasos manuales no documentados.","No ejecuta o faltan avances necesarios."),
    ("Síntesis escrita",6,"Conclusión dentro del notebook con pregunta, tipo, fuente, decisión y figura/tabla comprensible.","Conclusión escrita con evidencia incompleta.","Solo presenta números o no hay conclusión."),
    ("Límites y siguiente paso",4,"Dos límites concretos y una comprobación futura realizable sin causalidad indebida.","Límites vagos o siguiente paso poco claro.","Sin límites o hace afirmaciones causales no sustentadas."),
]
assert sum(x[1] for x in R1)==sum(x[1] for x in R2)==sum(x[1] for x in R3)==20
assert sum(x[1] for x in RF)==40


def first_task():
    doc=base(VERTICAL)
    title_and_meta(doc,"Proyecto Integrador · Tarea 1","Formativa · práctica guiada",4,"20 puntos","entrega al cierre del día")
    sections(doc,
        "Inicio del Proyecto Integrador: elige un dataset local y define un problema que se pueda responder con sus filas. Esta primera entrega audita los datos y decide el protocolo; aún no se califica el rendimiento de un modelo.",
        "Distinguir regresión, clasificación y clustering a partir de la respuesta buscada y justificar la pertinencia y calidad inicial del archivo elegido.",
        "Ficha y primeras celdas ejecutables de un único notebook que continuará durante las Tareas 2 y 3 y la práctica final. Debe existir una pregunta precisa, un tipo justificado, un diccionario breve y un plan de evaluación sin abrir Test.",
        ["En equipo pequeño o individualmente, escoge uno de los cinco CSV ya guardados en datos/ y mantén ese problema en las entregas sucesivas.",
         "Carga el archivo y explica qué representa una fila, su fuente, periodo, carácter público o sintético y si existe una variable objetivo.",
         "Muestra shape, head, dtypes, nulos, duplicados, categorías e identificadores; registra un extremo relevante y su posible efecto.",
         "Construye un primer gráfico con título, ejes y lectura; propón otros dos para la siguiente tarea.",
         "Esboza Train/Validation/Test conforme al objetivo y una referencia simple; no modeles usando Test."])
    option_table(doc)
    rubric(doc,R1,20,"Ficha y notebook inicial; Test aún reservado.")
    return doc


def second_task():
    doc=base(VERTICAL)
    title_and_meta(doc,"Proyecto Integrador · Tarea 2","Formativa · práctica guiada",5,"20 puntos","entrega al cierre del día")
    sections(doc,
        "Continúa el mismo dataset y la pregunta de la Tarea 1. Esta entrega prepara filas y columnas de manera reproducible y permite entender qué información se modelará. PCA es una opción cuando aporta a la pregunta, no un requisito universal.",
        "Limpiar y visualizar los datos sin perder trazabilidad; dividirlos según la pregunta y convertir variables cualitativas pertinentes a entradas numéricas sin fuga de información.",
        "Notebook ampliado con tabla de decisiones de calidad, al menos tres figuras interpretadas, particiones Train/Validation/Test verificadas y preprocesamiento ajustado solo en Train.",
        ["Justifica nulos, duplicados, valores extremos y columnas indisponibles en el momento de predicción; indica cuántas filas se conservan.",
         "Presenta tres figuras distintas con título, ejes y una oración que responda qué muestran; no deduzcas causalidad.",
         "Divide antes de calcular medianas, escalas o categorías: estratifica clases si procede; respeta fechas en pronóstico y países completos si buscas generalizar a ellos.",
         "Identifica cualitativas desde el principio, pero ajusta OneHotEncoder/imputación/escalado dentro de Pipeline usando solo Train; aplica las reglas aprendidas a Validation y Test.",
         "En clustering no inventes y: separa datos para comprobar estabilidad/perfiles y escala usando solo Train. Si no hay categorías sustantivas, explica por qué One-Hot no aplica."])
    heading(doc,"Ejemplo de una decisión correcta")
    paragraph(doc,"En telco_churn_fallback.csv, Contract es una categoría y customerID un identificador. Separa clientes en Train/Validation/Test; luego aprende en Train columnas 0/1 para las categorías Contract con One-Hot. No asignes Month-to-month=1 y One year=2 como si fuesen magnitudes ni incluyas customerID por rutina.")
    rubric(doc,R2,20,"Mismo notebook, ahora con preparación y tres gráficos.")
    return doc


def third_task():
    doc=base(VERTICAL)
    title_and_meta(doc,"Proyecto Integrador · Tarea 3","Formativa · práctica guiada",6,"20 puntos","entrega al cierre del día")
    sections(doc,
        "Tercera etapa del mismo Proyecto Integrador. Sobre las particiones y transformaciones de la Tarea 2, construye referencias y candidatos; usa únicamente Train y Validation para seleccionar. Test sigue reservado.",
        "Comparar alternativas en condiciones justas, interpretar las métricas adecuadas al tipo de problema y dejar escrita una decisión antes de la evaluación final.",
        "Notebook con referencia, tabla de alternativas sobre las mismas filas de Validation, métrica/criterio justificados, selección congelada y dos límites concretos. No se exige Grid Search si no aporta.",
        ["Regresión: media de Train o persistencia como baseline y al menos un modelo con MAE/RMSE. Clasificación: mayoría como baseline y al menos dos clasificadores con matriz, Precision, Recall, F1 y Accuracy.",
         "Clustering: compara valores de k o K-Means/Ward mediante tamaños, Silhouette, centroides y estabilidad; no informes Accuracy sin etiquetas verdaderas.",
         "Ajusta candidatos y transformaciones sin fuga; compara sobre las mismas filas de Validation o CV apropiada.",
         "Explica qué error o perfil cambiaría una decisión real; reconoce costos, sensibilidad y limitaciones.",
         "Antes de mirar Test, registra modelo o k, columnas, parámetros, métrica prioritaria y dos límites. Esta decisión se verificará en la práctica final."])
    rubric(doc,R3,20,"Decisión registrada en el notebook; Test no consultado.")
    return doc


def final_task():
    doc=base(MANUAL)
    title_and_meta(doc,"Proyecto Integrador · Práctica final","Final · laboratorio integrador",7,"40 puntos","entrega escrita, sin exposición")
    sections(doc,
        "Cierra el mismo proyecto de las tres tareas guiadas. Reúne sus avances en un notebook reproducible y ejecuta la comprobación final del procedimiento que fijaste el Día 6. No se vuelve a calificar cada tarea ya entregada.",
        "Evaluar con honestidad el procedimiento retenido, interpretar el resultado frente a la referencia y comunicar por escrito qué se obtuvo, qué no se puede concluir y qué se comprobaría después.",
        "Un solo archivo .ipynb que se ejecuta de principio a fin con el CSV local, conserva la pregunta y los tres avances, muestra Test o estabilidad retenida una vez y termina en una síntesis ejecutiva escrita con una figura o tabla. No hay presentación oral ni diapositivas obligatorias.",
        ["Revisa la decisión registrada en la Tarea 3. No modifiques variables, k, hiperparámetros ni umbral para mejorar Test.",
         "Regresión y clasificación: aplica al Test el preprocesamiento aprendido sin reajustarlo, informa métricas pertinentes y compáralas con Validation y baseline.",
         "Clustering: interpreta tamaños, perfiles y estabilidad de filas retenidas. K-Means permite asignar filas nuevas; Ward no tiene predict y exige otra comprobación razonada.",
         "Integra en orden código, figuras y conclusiones de Tareas 1–3. Escribe una síntesis comprensible: pregunta, dataset, tipo, decisión, cifra o perfil central, dos límites y siguiente paso.",
         "Entrega el notebook; la comunicación se evalúa en su texto, no mediante exposición."])
    heading(doc,"Consistencia entre entregas")
    paragraph(doc,"Tarea 1 (20) + Tarea 2 (20) + Tarea 3 (20) = 60 puntos formativos ya calificados. Esta práctica final añade 40 puntos por comprobación, integración e interpretación nuevas; no vuelve a calificar cada gráfico o limpieza por segunda vez.")
    doc.add_page_break()  # La rúbrica empieza completa en una página interior.
    rubric(doc,RF,40,"Notebook integrado y síntesis escrita, sin exposición oral.")
    return doc


OUTPUTS = [
    ("Instrumento_01_Diagnostica.docx", diagnostic, VERTICAL),
    ("Instrumento_02_Proyecto_Tarea1.docx", first_task, VERTICAL),
    ("Instrumento_03_Proyecto_Tarea2.docx", second_task, VERTICAL),
    ("Instrumento_04_Proyecto_Tarea3.docx", third_task, VERTICAL),
    ("Instrumento_05_Proyecto_Practica_Final.docx", final_task, MANUAL),
]


def customize_cover(path):
    temp=path.with_name(path.stem+".temporal.docx")
    replacements={
        "Nombre del Curso":"Proyecto Integrador",
        "Tema de la Unidad":"Práctica final · Módulo 3",
        "Expositor":"Diplomado Big Data e IA",
    }
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(temp,"w") as target:
        for info in source.infolist():
            content=source.read(info.filename)
            if info.filename=="word/document.xml":
                xml=content.decode("utf-8")
                for original,new in replacements.items():
                    old=f">{original}<"
                    assert old in xml,original
                    xml=xml.replace(old,">"+html.escape(new)+"<")
                content=xml.encode("utf-8")
            target.writestr(info,content)
    os.replace(temp,path)


if __name__=="__main__":
    assert VERTICAL.is_file() and MANUAL.is_file()
    for name,builder,template in OUTPUTS:
        doc=builder()
        output=ROOT/name
        doc.save(output)
        if template==MANUAL:customize_cover(output)
        print(output.name,output.stat().st_size,"bytes", "base",template.name)
